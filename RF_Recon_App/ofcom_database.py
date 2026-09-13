import sqlite3
import os
import time
import openpyxl

class OfcomDatabaseManager:
    def __init__(self, db_path="ofcom_tv.db"):
        self.db_path = db_path
        self.last_update_key = "last_update"
        self.cache_days = 30
        
        # Initialize DB
        self.conn = sqlite3.connect(self.db_path, check_same_thread=False)
        self._create_tables()

    def _create_tables(self):
        cursor = self.conn.cursor()
        # Create stations table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS stations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                site_name TEXT,
                lat REAL,
                lon REAL,
                channel INTEGER,
                erp REAL,
                multiplex TEXT
            )
        ''')
        # Create metadata table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS metadata (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        ''')
        self.conn.commit()
        
    def get_last_update_time(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT value FROM metadata WHERE key = ?", (self.last_update_key,))
        row = cursor.fetchone()
        if row:
            try:
                return float(row[0])
            except ValueError:
                return 0.0
        return 0.0

    def needs_update(self):
        last_update = self.get_last_update_time()
        if last_update == 0.0:
            return True
        days_since_update = (time.time() - last_update) / (24 * 3600)
        return days_since_update > self.cache_days
        
    def has_data(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM stations")
        return cursor.fetchone()[0] > 0

    def import_excel(self, file_path, progress_callback=None):
        """
        Parses the Ofcom XLSX file.
        The Ofcom TV transmitter data has sheets like 'PSB_COM_Muxes' and 'LTVMux_NIMux_GIMux'.
        """
        if progress_callback: progress_callback(10, "Loading Excel File...")
        
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            if progress_callback: progress_callback(100, f"Error: {str(e)}")
            return False
            
        stations = []
        
        for sheet_name in wb.sheetnames:
            if progress_callback: progress_callback(50, f"Parsing Sheet: {sheet_name}")
            sheet = wb[sheet_name]
            
            # Find header row
            header_row_idx = None
            headers = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if row and any(isinstance(c, str) and 'Site' in c for c in row):
                    header_row_idx = i
                    headers = [str(c).strip().lower() if c else "" for c in row]
                    break
            
            if header_row_idx is None:
                continue
                
            site_idx = next((i for i, h in enumerate(headers) if 'site' in h), -1)
            lat_idx = next((i for i, h in enumerate(headers) if 'lat' in h), -1)
            lon_idx = next((i for i, h in enumerate(headers) if 'long' in h), -1)
            chan_idx = next((i for i, h in enumerate(headers) if 'channel' in h or 'uhf' in h), -1)
            erp_idx = next((i for i, h in enumerate(headers) if 'erp' in h), -1)
            mux_idx = next((i for i, h in enumerate(headers) if 'mux' in h or 'multiplex' in h), -1)
            
            if site_idx == -1 or chan_idx == -1:
                continue
                
            for row in sheet.iter_rows(min_row=header_row_idx+2, values_only=True):
                site = row[site_idx]
                chan = row[chan_idx]
                if not site or not chan:
                    continue
                    
                lat = row[lat_idx] if lat_idx != -1 else 0.0
                lon = row[lon_idx] if lon_idx != -1 else 0.0
                erp = row[erp_idx] if erp_idx != -1 else 0.0
                mux = row[mux_idx] if mux_idx != -1 else ""
                
                try:
                    chan = int(chan)
                    lat = float(lat) if lat else 0.0
                    lon = float(lon) if lon else 0.0
                    erp = float(erp) if erp else 0.0
                except (ValueError, TypeError):
                    continue
                    
                stations.append((str(site), lat, lon, chan, erp, str(mux)))
                
        if progress_callback: progress_callback(90, "Saving to Database...")
        cursor = self.conn.cursor()
        cursor.execute("DELETE FROM stations")
        cursor.executemany("INSERT INTO stations (site_name, lat, lon, channel, erp, multiplex) VALUES (?, ?, ?, ?, ?, ?)", stations)
        cursor.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES (?, ?)", (self.last_update_key, str(time.time())))
        self.conn.commit()
        
        if progress_callback: progress_callback(100, "Import Complete")
        return True
