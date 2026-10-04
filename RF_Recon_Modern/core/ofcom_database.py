import sqlite3
try:
    from .app_paths import user_database
except ImportError:
    from app_paths import user_database
import threading
import math
import os
import time
import openpyxl


def _haversine_km(lat1, lon1, lat2, lon2):
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = (math.sin(dlat / 2) ** 2 +
         math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(dlon / 2) ** 2)
    return 6371.0 * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


# --- Grid references to latitude / longitude -----------------------------------
#
# Ofcom gives transmitter sites as Ordnance Survey grid references (Great Britain)
# or Irish Grid references (Northern Ireland), not latitude and longitude.

def _tm_to_latlon(e, n, a, b, f0, lat0, lon0, e0, n0):
    """Transverse Mercator easting/northing to latitude/longitude (radians) on its own ellipsoid."""
    e2 = 1 - (b * b) / (a * a)
    nn = (a - b) / (a + b)
    lat = lat0
    m = 0.0
    while True:
        lat = (n - n0 - m) / (a * f0) + lat
        ma = (1 + nn + 1.25 * nn ** 2 + 1.25 * nn ** 3) * (lat - lat0)
        mb = (3 * nn + 3 * nn ** 2 + 2.625 * nn ** 3) * math.sin(lat - lat0) * math.cos(lat + lat0)
        mc = (1.875 * nn ** 2 + 1.875 * nn ** 3) * math.sin(2 * (lat - lat0)) * math.cos(2 * (lat + lat0))
        md = (35 / 24) * nn ** 3 * math.sin(3 * (lat - lat0)) * math.cos(3 * (lat + lat0))
        m = b * f0 * (ma - mb + mc - md)
        if abs(n - n0 - m) < 1e-5:
            break
    sin_lat, cos_lat, tan_lat = math.sin(lat), math.cos(lat), math.tan(lat)
    nu = a * f0 / math.sqrt(1 - e2 * sin_lat ** 2)
    rho = a * f0 * (1 - e2) / (1 - e2 * sin_lat ** 2) ** 1.5
    eta2 = nu / rho - 1
    de = e - e0
    vii = tan_lat / (2 * rho * nu)
    viii = tan_lat / (24 * rho * nu ** 3) * (5 + 3 * tan_lat ** 2 + eta2 - 9 * tan_lat ** 2 * eta2)
    ix = tan_lat / (720 * rho * nu ** 5) * (61 + 90 * tan_lat ** 2 + 45 * tan_lat ** 4)
    x = 1 / (cos_lat * nu)
    xi = 1 / (6 * cos_lat * nu ** 3) * (nu / rho + 2 * tan_lat ** 2)
    xii = 1 / (120 * cos_lat * nu ** 5) * (5 + 28 * tan_lat ** 2 + 24 * tan_lat ** 4)
    return (lat - vii * de ** 2 + viii * de ** 4 - ix * de ** 6,
            lon0 + x * de - xi * de ** 3 + xii * de ** 5)


def _to_wgs84(lat, lon, a, b, tx, ty, tz, rx, ry, rz, s_ppm):
    """A latitude/longitude on a local datum to WGS84 degrees (Helmert transformation)."""
    e2 = 1 - (b * b) / (a * a)
    nu = a / math.sqrt(1 - e2 * math.sin(lat) ** 2)
    x = nu * math.cos(lat) * math.cos(lon)
    y = nu * math.cos(lat) * math.sin(lon)
    z = (1 - e2) * nu * math.sin(lat)
    s = 1 + s_ppm * 1e-6
    rx, ry, rz = (math.radians(v / 3600) for v in (rx, ry, rz))
    x2 = tx + s * x - rz * y + ry * z
    y2 = ty + rz * x + s * y - rx * z
    z2 = tz - ry * x + rx * y + s * z
    wa, wb = 6378137.0, 6356752.3141
    we2 = 1 - (wb * wb) / (wa * wa)
    p = math.hypot(x2, y2)
    lat2 = math.atan2(z2, p * (1 - we2))
    for _ in range(6):
        nu2 = wa / math.sqrt(1 - we2 * math.sin(lat2) ** 2)
        lat2 = math.atan2(z2 + we2 * nu2 * math.sin(lat2), p)
    return math.degrees(lat2), math.degrees(math.atan2(y2, x2))


def osgb_to_latlon(easting, northing):
    """Ordnance Survey National Grid (OSGB36) metres to WGS84 degrees, good to a few metres."""
    a, b = 6377563.396, 6356256.909                       # Airy 1830
    lat, lon = _tm_to_latlon(easting, northing, a, b, 0.9996012717, math.radians(49), math.radians(-2), 400000, -100000)
    return _to_wgs84(lat, lon, a, b, 446.448, -125.157, 542.060, 0.1502, 0.2470, 0.8421, -20.4894)


def irish_to_latlon(easting, northing):
    """Irish Grid metres to WGS84 degrees, good to a few metres."""
    a, b = 6377340.189, 6356034.447                       # Airy modified
    lat, lon = _tm_to_latlon(easting, northing, a, b, 1.000035, math.radians(53.5), math.radians(-8), 200000, 250000)
    return _to_wgs84(lat, lon, a, b, 482.53, -130.596, 564.557, -1.042, -0.214, -0.631, 8.15)


def _grid_ref_metres(ref, irish=False):
    """'TL20474944' (GB, two letters) or 'J302749' (Irish, one letter) to (easting, northing), or None."""
    ref = "".join(str(ref or "").upper().split())
    n_letters = 1 if irish else 2
    letters, digits = ref[:n_letters], ref[n_letters:]
    if len(letters) != n_letters or not letters.isalpha() or not digits.isdigit() or len(digits) % 2 or not digits:
        return None
    half = len(digits) // 2
    scale = 10 ** (5 - half)
    de, dn = int(digits[:half]) * scale, int(digits[half:]) * scale

    def idx(ch):                                           # A-Z without I, on a 5 x 5 grid
        i = ord(ch) - ord("A")
        return i - 1 if i > 7 else i

    if irish:
        i = idx(letters)
        return (i % 5) * 100000 + de, (4 - i // 5) * 100000 + dn
    i1, i2 = idx(letters[0]), idx(letters[1])
    e100 = ((i1 - 2) % 5) * 5 + (i2 % 5)
    n100 = 19 - (i1 // 5) * 5 - (i2 // 5)
    return e100 * 100000 + de, n100 * 100000 + dn


def _site_position(ngr, ngr_x, ngr_y, irish, irish_x, irish_y, lat=None, lon=None):
    """
    (lat, lon) of a site from whatever the row gives. The grid reference text is
    preferred over the separate easting/northing columns: Ofcom's sheet has rows
    whose numeric columns are wrong (a northing of 0) while the reference is right.
    """
    def num(v):
        try:
            return float(str(v).strip())
        except (TypeError, ValueError):
            return None
    la, lo = num(lat), num(lon)
    if la and lo:
        return la, lo
    en = _grid_ref_metres(ngr)
    if en is None and num(ngr_x) and num(ngr_y):
        en = (num(ngr_x), num(ngr_y))
    if en is not None:
        return osgb_to_latlon(*en)
    en = _grid_ref_metres(irish, irish=True)
    if en is None and num(irish_x) and num(irish_y):
        en = (num(irish_x), num(irish_y))
    if en is not None:
        return irish_to_latlon(*en)
    return None


class OfcomDatabaseManager:
    def __init__(self, db_path="ofcom_tv.db"):
        # Relative names refer to the per-user copy (writable, independent of the
        # working directory); absolute paths are used as given.
        if not os.path.isabs(db_path):
            db_path = user_database(db_path)
        self.db_path = db_path
        self.last_update_key = "last_update"
        self.cache_days = 30
        
        # Initialize DB (connection shared across threads; every use holds _db_lock)
        self._db_lock = threading.RLock()
        self.conn = sqlite3.connect(self.db_path, check_same_thread=False)
        self._create_tables()
        self._seed_from_bundled()

    # fReqon ships with Ofcom's transmitter data (ofcom_tv.db in the application folder,
    # built from Ofcom's spreadsheet by tools/build_ofcom_db.py). The user's copy starts
    # from it, takes a newer shipped one when the application is updated, and is
    # replaced by whatever spreadsheet the user imports.
    def _seed_from_bundled(self):
        from .app_paths import APP_DIR
        bundled = APP_DIR / "ofcom_tv.db"
        try:
            if not bundled.exists() or os.path.samefile(bundled, self.db_path):
                return
        except OSError:
            return
        mine = self.info()
        if self.has_data() and mine["source"] == "import":
            return                                  # the user's own import stands
        # ("legacy": imported before the importer could read Ofcom's grid references, so
        # every site is at 0, 0 and no lookup matches; it is replaced)
        try:
            src = sqlite3.connect(f"file:{bundled}?mode=ro", uri=True)
            rows = src.execute("SELECT site_name, lat, lon, channel, erp, multiplex FROM stations").fetchall()
            meta = dict(src.execute("SELECT key, value FROM metadata").fetchall())
            src.close()
        except sqlite3.Error:
            return
        if not rows or (self.has_data() and mine["source"] == "bundled"
                        and (meta.get("data_date") or "") <= (mine["data_date"] or "")):
            return                                  # nothing shipped, or nothing newer than what is here
        with self._db_lock, self.conn:
            self.conn.execute("DELETE FROM stations")
            self.conn.executemany("INSERT INTO stations (site_name, lat, lon, channel, erp, multiplex) VALUES (?, ?, ?, ?, ?, ?)", rows)
            for key, value in dict(meta, source="bundled").items():
                self.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES (?, ?)", (key, value))

    def info(self) -> dict:
        """{"source": "bundled" | "import" | "legacy" (an import from before positions could be
        read) | None, "data_date": the spreadsheet's own date (ISO) or None, "records"}"""
        meta = dict(self._query("SELECT key, value FROM metadata"))
        n = self._query("SELECT COUNT(*) FROM stations")[0][0]
        source = meta.get("source") or ("legacy" if n else None)
        return {"source": source if n else None, "data_date": meta.get("data_date"), "records": n}

    def _create_tables(self):
        cursor = self.conn.cursor()
        # Create stations table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS stations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                site_name TEXT,
                lat REAL,
                lon REAL,
                channel INTEGER,
                erp REAL,
                multiplex TEXT
            )
        ''')
        # Create metadata table
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS metadata (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        ''')
        self.conn.commit()
        
    def _query(self, sql, params=()):
        with self._db_lock:
            return self.conn.execute(sql, params).fetchall()

    def get_last_update_time(self):
        rows = self._query("SELECT value FROM metadata WHERE key = ?", (self.last_update_key,))
        row = rows[0] if rows else None
        if row:
            try:
                return float(row[0])
            except ValueError:
                return 0.0
        return 0.0

    def needs_update(self):
        last_update = self.get_last_update_time()
        if last_update == 0.0:
            return True
        days_since_update = (time.time() - last_update) / (24 * 3600)
        return days_since_update > self.cache_days
        
    def has_data(self):
        return self._query("SELECT COUNT(*) FROM stations")[0][0] > 0

    @staticmethod
    def _geocode_postcode(postcode: str):
        """Return (lat, lon) for a UK postcode, or None. Tries the full postcode,
        then the outward code (e.g. 'SW1A' from 'SW1A 1AA')."""
        import pgeocode
        code = " ".join(postcode.upper().split())
        if " " in code:
            outward = code.split(" ")[0]
        elif len(code) > 4:
            outward = code[:-3]  # inward code is always 3 characters
        else:
            outward = code
        attempts = [("GB_full", code), ("GB", outward)]
        for country, query in attempts:
            try:
                res = pgeocode.Nominatim(country).query_postal_code(query)
                if not (math.isnan(res.latitude) or math.isnan(res.longitude)):
                    return float(res.latitude), float(res.longitude)
            except Exception:
                continue
        return None

    def query_by_postcode(self, postcode: str, radius_km: float = 80):
        """
        Returns imported Ofcom transmitter sites within radius_km of a UK postcode,
        nearest first, in the same shape as FCCDatabaseManager.get_stations_near_zip.
        """
        loc = self._geocode_postcode(postcode)
        if loc is None:
            return []
        target_lat, target_lon = loc

        lat_delta = radius_km / 111.0
        lon_delta = radius_km / (111.0 * max(0.1, math.cos(math.radians(target_lat))))
        rows = self._query(
            "SELECT site_name, lat, lon, channel, erp, multiplex FROM stations "
            "WHERE lat BETWEEN ? AND ? AND lon BETWEEN ? AND ?",
            (target_lat - lat_delta, target_lat + lat_delta, target_lon - lon_delta, target_lon + lon_delta)
        )
        stations = []
        for site, lat, lon, channel, erp, mux in rows:
            if not lat or not lon:
                continue
            dist = _haversine_km(target_lat, target_lon, lat, lon)
            if dist <= radius_km:
                stations.append({
                    "site_name": site,
                    "call_sign": f"{site} ({mux})" if mux else site,
                    "distance_km": round(dist, 1),
                    "channel": channel,
                    "erp": erp,
                    "is_public_safety": False,
                })
        stations.sort(key=lambda s: s["distance_km"])
        return stations

    def import_excel(self, file_path, progress_callback=None, source="import"):
        """
        Parses the Ofcom XLSX file.
        The Ofcom TV transmitter data has sheets like 'PSB_COM_Muxes' and 'LTVMux_NIMux_GIMux'.
        progress_callback runs on the calling thread; if this is called from a
        worker thread, pass a signal's emit rather than a widget-updating function.

        Returns the number of transmitter records imported. Raises ValueError if
        the file can't be read or contains no recognisable records, in which
        case the existing data is left untouched.
        """
        if progress_callback: progress_callback(10, "Loading Excel File...")

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise ValueError(f"Could not open the spreadsheet: {e}") from e
            
        stations = []
        skipped_sites = 0

        for sheet_name in wb.sheetnames:
            if progress_callback: progress_callback(50, f"Parsing Sheet: {sheet_name}")
            sheet = wb[sheet_name]
            
            rows = list(sheet.iter_rows(values_only=True))
            header_row_idx = next((i for i, row in enumerate(rows)
                                   if row and any(isinstance(c, str) and c.strip().lower() == "site name" for c in row)), None)
            if header_row_idx is None:
                continue
            headers = [str(c).strip().lower() if c else "" for c in rows[header_row_idx]]

            def col(*names):
                return next((i for i, h in enumerate(headers) if h in names), -1)

            site_idx = col("site name")
            ngr_idx, ngr_x_idx, ngr_y_idx = col("gbngr"), col("gbngrx"), col("gbngry")
            irish_idx, irish_x_idx, irish_y_idx = col("irishgr"), col("irishgrx"), col("irishgry")
            lat_idx = next((i for i, h in enumerate(headers) if h.startswith("lat")), -1)
            lon_idx = next((i for i, h in enumerate(headers) if h.startswith("long")), -1)
            # One row is one site with several multiplexes: every "<mux> channel" column,
            # with the ERP column of the same multiplex ("PSB1 Channel" / "PSB1 ERP (kW)",
            # "LTVMux Channel" / "LTV ERP (kW)")
            muxes = []
            for i, h in enumerate(headers):
                if h.endswith(" channel"):
                    name = rows[header_row_idx][i].strip()[:-len(" channel")]
                    key = name.lower().replace("mux", "")
                    erp_i = next((j for j, hj in enumerate(headers) if "erp" in hj and hj.replace("mux", "").startswith(key)), -1)
                    muxes.append((name, i, erp_i))
            if site_idx == -1 or not muxes:
                continue

            for row in rows[header_row_idx + 1:]:
                # Rows can be shorter than the header (trailing empty cells)
                def cell(idx, default=None):
                    return row[idx] if 0 <= idx < len(row) else default

                site = cell(site_idx)
                if not site:
                    continue
                pos = _site_position(cell(ngr_idx), cell(ngr_x_idx), cell(ngr_y_idx),
                                     cell(irish_idx), cell(irish_x_idx), cell(irish_y_idx),
                                     cell(lat_idx), cell(lon_idx))
                if pos is None:
                    skipped_sites += 1
                    continue
                for mux, chan_i, erp_i in muxes:
                    try:
                        chan = int(float(str(cell(chan_i)).strip()))
                    except (ValueError, TypeError):
                        continue                    # this site does not carry that multiplex
                    try:
                        erp = float(str(cell(erp_i)).strip()) if cell(erp_i) not in (None, "") else 0.0
                    except (ValueError, TypeError):
                        erp = 0.0
                    stations.append((str(site).strip(), pos[0], pos[1], chan, erp, mux))

        if not stations:
            raise ValueError("No transmitter records found. Is this Ofcom's television "
                             "transmitter frequency spreadsheet?")
        if progress_callback: progress_callback(90, "Saving to Database...")
        with self._db_lock, self.conn:  # atomic: commits on success, rolls back on error
            self.conn.execute("DELETE FROM stations")
            self.conn.executemany("INSERT INTO stations (site_name, lat, lon, channel, erp, multiplex) VALUES (?, ?, ?, ?, ?, ?)", stations)
            self.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES (?, ?)", (self.last_update_key, str(time.time())))
            # The spreadsheet's own date (when Ofcom last saved it) says how current the data is
            stamp = getattr(wb.properties, "modified", None) or getattr(wb.properties, "created", None)
            self.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES ('data_date', ?)",
                              (stamp.strftime("%Y-%m-%d") if stamp else "",))
            self.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES ('source', ?)", (source,))
        
        if progress_callback: progress_callback(100, "Import Complete")
        return len(stations)
