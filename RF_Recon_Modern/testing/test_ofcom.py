"""
test_ofcom.py - Ofcom transmitter spreadsheet import: grid references to
latitude/longitude and the multi-multiplex layout (no network, no analyzer).

    python testing/test_ofcom.py
"""

import os
import sys
import tempfile
import time

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from core.ofcom_database import OfcomDatabaseManager, osgb_to_latlon, irish_to_latlon, _grid_ref_metres, _site_position


def test_grid_references():
    assert _grid_ref_metres("TQ33937122") == (533930, 171220)            # Crystal Palace
    lat, lon = osgb_to_latlon(533930, 171220)
    assert abs(lat - 51.4242) < 0.0005 and abs(lon - -0.0750) < 0.0005, (lat, lon)
    assert _grid_ref_metres("SD 660 144") == (366000, 414400)            # 6 figures, with spaces (Winter Hill)
    assert _grid_ref_metres("J28677512", irish=True) == (328670, 375120)  # Divis
    lat, lon = irish_to_latlon(328670, 375120)
    assert abs(lat - 54.6079) < 0.001 and abs(lon - -6.0094) < 0.001, (lat, lon)
    assert _grid_ref_metres("") is None and _grid_ref_metres("nonsense") is None and _grid_ref_metres("TQ123") is None
    # The reference text wins over wrong numeric columns (Ofcom's Sandy Heath row has a northing of 0)
    lat, lon = _site_position("TL20474944", "520470", "0", None, None, None)
    assert abs(lat - 52.130) < 0.001 and abs(lon - -0.2415) < 0.001, (lat, lon)
    assert _site_position(None, "520470", "249440", None, None, None) is not None       # numbers alone are used
    assert _site_position(None, None, None, None, None, None) is None


def _workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PSB_COM_Muxes"
    ws.append(["ITV Region", "Transmitter Group", "Site Number", "Site Name", "GBNGR", "GBNGRX", "GBNGRY",
               "IrishGR", "IrishGRX", "IrishGRY", "PSB1 Channel", "PSB1 Offset", "PSB1 ERP (kW)",
               "PSB2 Channel", "PSB2 Offset", "PSB2 ERP (kW)", "COM4 Channel", "COM4 Offset", "COM4 ERP (kW)", "Notes"])
    ws.append(["London", "Crystal Palace", "10100", "Crystal Palace", "TQ33937122", "533930", "171220", None, None, None,
               "23", None, "200", "26", None, "200", "25", None, "200", None])
    ws.append(["London", "Crystal Palace", "10101", "Small Relay", "TQ30008000", "530000", "180000", None, None, None,
               "41", "+", "0.01", "44", None, "0.01", None, None, None, None])           # carries no COM4
    ws.append(["Ulster", "Divis", "30100", "Divis", None, None, None, "J28677512", "328670", "375120",
               "27", None, "100", "21", None, "100", "23", None, "100", None])
    ws.append(["x", "y", "1", "Nowhere", None, None, None, None, None, None, "30", None, "1", None, None, None, None, None, None, None])
    ws2 = wb.create_sheet("LTVMux_NIMux_GIMux")
    ws2.append(["ITV Region", "Transmitter Group", "Site Number", "Site Name", "GBNGR", "GBNGRX", "GBNGRY",
                "IrishGR", "IrishGRX", "IrishGRY", "LTVMux Channel", "LTVMux Offset", "LTV ERP (kW)", "NIMux Channel", "NIMux ERP (kW)"])
    ws2.append(["London", "Crystal Palace", "10100", "Crystal Palace", "TQ33937122", "533930", "171220", None, None, None,
                "35", None, "20", None, None])
    wb.save(path)


def test_import():
    d = tempfile.mkdtemp()
    xlsx = os.path.join(d, "ofcom.xlsx")
    _workbook(xlsx)
    db = OfcomDatabaseManager(os.path.join(d, "o.db"))
    assert db.import_excel(xlsx) == 9                 # 3 + 2 + 3 + 1; the site with no position is left out
    rows = db._query("SELECT site_name, round(lat, 3), round(lon, 3), channel, erp, multiplex FROM stations ORDER BY id")
    assert rows[0] == ("Crystal Palace", 51.424, -0.075, 23, 200.0, "PSB1") and rows[2][3:] == (25, 200.0, "COM4"), rows[:3]
    assert [r[3] for r in rows if r[0] == "Small Relay"] == [41, 44]
    assert [(r[1], r[2]) for r in rows if r[0] == "Divis"][0] == (54.608, -6.009)
    assert rows[-1] == ("Crystal Palace", 51.424, -0.075, 35, 20.0, "LTVMux")
    assert db.has_data()
    # A lookup near Crystal Palace finds its channels with their powers (the geocoder is stubbed: no network)
    OfcomDatabaseManager._geocode_postcode = staticmethod(lambda pc: (51.42, -0.08))
    found = db.query_by_postcode("SE19 1UE", 30)
    assert sorted(s["channel"] for s in found) == [23, 25, 26, 35, 41, 44]
    cp = next(s for s in found if s["channel"] == 23)
    assert cp["erp"] == 200.0 and cp["distance_km"] < 1.0 and "PSB1" in cp["call_sign"]
    # Something that is not Ofcom's sheet is refused and the data kept
    bad = os.path.join(d, "bad.xlsx")
    wb = openpyxl.Workbook(); wb.active.append(["a", "b"]); wb.save(bad)
    try:
        db.import_excel(bad)
        raise AssertionError("an unrelated spreadsheet must be refused")
    except ValueError:
        pass
    assert db.has_data()


def test_shipped_data():
    """fReqon ships with Ofcom's data: a new user copy starts from it, the user's own import stands."""
    d = tempfile.mkdtemp()
    db = OfcomDatabaseManager(os.path.join(d, "user.db"))
    info = db.info()
    assert info["source"] == "bundled" and info["records"] > 3000 and info["data_date"] >= "2024-10-08", info
    assert db._query("SELECT COUNT(DISTINCT site_name) FROM stations")[0][0] > 1000
    cp = db._query("SELECT lat, lon, erp FROM stations WHERE site_name = 'Crystal Palace' AND multiplex = 'PSB1'")
    assert len(cp) == 1 and abs(cp[0][0] - 51.424) < 0.005 and abs(cp[0][1] - -0.075) < 0.005 and cp[0][2] == 200.0, cp
    # The user's import replaces it, and is kept the next time the app starts
    xlsx = os.path.join(d, "ofcom.xlsx")
    _workbook(xlsx)
    assert db.import_excel(xlsx) == 9 and db.info()["source"] == "import"
    db.conn.close()
    db = OfcomDatabaseManager(os.path.join(d, "user.db"))
    assert db.info()["source"] == "import" and db.info()["records"] == 9
    # An import from before the importer could read positions (no source recorded) is replaced
    with db._db_lock, db.conn:
        db.conn.execute("DELETE FROM metadata WHERE key IN ('source', 'data_date')")
    db.conn.close()
    db = OfcomDatabaseManager(os.path.join(d, "user.db"))
    assert db.info()["source"] == "bundled" and db.info()["records"] > 3000
    # A copy seeded from an older shipped file takes the newer shipped one
    with db._db_lock, db.conn:
        db.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES ('source', 'bundled')")
        db.conn.execute("INSERT OR REPLACE INTO metadata (key, value) VALUES ('data_date', '2020-01-01')")
    db.conn.close()
    db = OfcomDatabaseManager(os.path.join(d, "user.db"))
    assert db.info()["records"] > 3000 and db.info()["source"] == "bundled"


if __name__ == "__main__":
    for test in (test_grid_references, test_import, test_shipped_data):
        t0 = time.monotonic()
        test()
        print(f"ok  {test.__name__}  ({time.monotonic() - t0:.1f} s)")
    print("All Ofcom tests passed.")
